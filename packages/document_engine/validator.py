"""
ARM Stage 9 - Deterministic Document Validator
Provides structural validation for generated DOCX OpenXML packages,
reopen verification, and SHA-256 template immutability auditing.
"""

import os
import zipfile
import hashlib
from typing import Dict, Any, Optional
import docx

from apps.api.core.logging import get_logger

logger = get_logger("document.validator")


class DocumentValidator:
    """
    Validates generated DOCX documents and verifies template file immutability.
    """

    @staticmethod
    def calculate_sha256(file_path_or_bytes) -> str:
        """Calculates SHA-256 checksum of a file on disk or bytes."""
        sha = hashlib.sha256()
        if isinstance(file_path_or_bytes, bytes):
            sha.update(file_path_or_bytes)
            return sha.hexdigest()
        if not os.path.exists(file_path_or_bytes):
            raise FileNotFoundError(f"File not found for SHA-256 calculation: {file_path_or_bytes}")
        with open(file_path_or_bytes, "rb") as f:
            while chunk := f.read(65536):
                sha.update(chunk)
        return sha.hexdigest()

    @staticmethod
    def verify_template_immutability(template_path: str, expected_sha256: str) -> bool:
        """
        Verifies that an original template file has not been mutated.
        Returns True if hash matches, False otherwise.
        """
        current_sha = DocumentValidator.calculate_sha256(template_path)
        is_identical = (current_sha == expected_sha256)
        if not is_identical:
            logger.error(
                f"Template immutability violation! Expected {expected_sha256}, got {current_sha}"
            )
        return is_identical

    @staticmethod
    def validate_docx(file_path: str) -> Dict[str, Any]:
        """
        Performs comprehensive structural validation and reopen check on a generated DOCX file:
        1. Checks file existence and non-empty byte size
        2. Validates ZIP package integrity
        3. Verifies essential OpenXML parts: [Content_Types].xml, _rels/.rels, word/document.xml
        4. Reopens document with python-docx and verifies structural elements
        """
        result: Dict[str, Any] = {
            "valid": False,
            "reopen_success": False,
            "file_size_bytes": 0,
            "paragraph_count": 0,
            "table_count": 0,
            "section_count": 0,
            "word_count": 0,
            "errors": [],
            "warnings": [],
        }

        # 1. Existence and size check
        if not os.path.exists(file_path):
            result["errors"].append(f"Generated DOCX file does not exist at path: {file_path}")
            return result

        size = os.path.getsize(file_path)
        result["file_size_bytes"] = size
        if size == 0:
            result["errors"].append("Generated DOCX file is 0 bytes.")
            return result

        # 2. ZIP package validation
        if not zipfile.is_zipfile(file_path):
            result["errors"].append("Generated file is not a valid ZIP package (DOCX container corrupted).")
            return result

        # 3. OpenXML structure check
        try:
            with zipfile.ZipFile(file_path, "r") as zf:
                namelist = zf.namelist()
                required_parts = [
                    "[Content_Types].xml",
                    "_rels/.rels",
                    "word/document.xml"
                ]
                for part in required_parts:
                    if part not in namelist:
                        result["errors"].append(f"Missing mandatory OpenXML part: {part}")

                # Test zip integrity
                bad_file = zf.testzip()
                if bad_file:
                    result["errors"].append(f"Corrupted file inside DOCX package: {bad_file}")
        except Exception as e:
            result["errors"].append(f"Error inspecting DOCX OpenXML ZIP contents: {e}")
            return result

        if result["errors"]:
            return result

        # 4. Mandatory Reopen Test via python-docx
        try:
            doc = docx.Document(file_path)
            result["section_count"] = len(doc.sections)
            result["paragraph_count"] = len(doc.paragraphs)
            result["table_count"] = len(doc.tables)

            # Calculate word count across all paragraphs and table cells
            words = 0
            for p in doc.paragraphs:
                words += len(p.text.split())
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        words += len(cell.text.split())
            result["word_count"] = words

            result["reopen_success"] = True
            result["valid"] = True
            result["is_valid"] = True
            logger.info(
                f"DOCX validation PASSED: {result['section_count']} sections, "
                f"{result['paragraph_count']} paragraphs, {result['table_count']} tables, "
                f"{result['word_count']} words."
            )
        except Exception as e:
            result["reopen_success"] = False
            result["valid"] = False
            result["is_valid"] = False
            result["errors"].append(f"Reopen test failed with python-docx: {e}")
            logger.error(f"DOCX Reopen test FAILED for {file_path}: {e}")

        result["is_valid"] = result["valid"]
        return result
