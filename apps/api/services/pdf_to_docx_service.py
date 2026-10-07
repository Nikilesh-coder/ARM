"""
ARM — PDF to DOCX Template Conversion & Validation Service
Converts uploaded institutional PDF templates into editable DOCX working representations
while preserving the original uploaded PDF as the immutable source of truth.
Prioritizes text positioning, tables, images, layout, margins, and headers/footers.
"""

import os
import io
import tempfile
from typing import Tuple, Dict, Any, Optional
import docx
from pdf2docx import Converter

from apps.api.core.logging import get_logger
from apps.api.services.pdf_converter import pdf_converter_registry

logger = get_logger("service.pdf_to_docx")

CONVERSION_ERROR_MESSAGE = (
    "We couldn't convert this PDF into an editable template. "
    "Please try another PDF or upload a DOCX template."
)


class PdfToDocxService:
    """Production service for converting and validating PDF templates into DOCX working copies."""

    @classmethod
    def validate_converted_docx(cls, docx_path: str) -> Dict[str, Any]:
        """
        Validates that the converted DOCX:
        1. Exists and is non-empty.
        2. Can be parsed without errors by python-docx.
        3. Has valid paragraphs, tables, or sections.
        4. Contains meaningful text or content.
        """
        if not os.path.exists(docx_path):
            raise ValueError(CONVERSION_ERROR_MESSAGE)

        file_size = os.path.getsize(docx_path)
        if file_size == 0:
            raise ValueError(CONVERSION_ERROR_MESSAGE)

        try:
            doc = docx.Document(docx_path)
        except Exception as e:
            logger.error(f"[PDF-CONVERSION-VALIDATION] Failed to open generated DOCX: {e}")
            raise ValueError(CONVERSION_ERROR_MESSAGE)

        sections_count = len(doc.sections)
        paragraphs_count = len(doc.paragraphs)
        tables_count = len(doc.tables)

        # Count text characters across paragraphs and tables
        para_chars = sum(len(p.text.strip()) for p in doc.paragraphs)
        table_chars = 0
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_chars += len(cell.text.strip())

        total_chars = para_chars + table_chars

        if sections_count < 1 or (paragraphs_count == 0 and tables_count == 0):
            logger.warning("[PDF-CONVERSION-VALIDATION] Converted document has no paragraphs or tables.")
            raise ValueError(CONVERSION_ERROR_MESSAGE)

        if total_chars < 5 and tables_count == 0:
            logger.warning(f"[PDF-CONVERSION-VALIDATION] Converted document has insufficient text content ({total_chars} chars).")
            raise ValueError(CONVERSION_ERROR_MESSAGE)

        stats = {
            "docx_path": docx_path,
            "file_size": file_size,
            "sections_count": sections_count,
            "paragraphs_count": paragraphs_count,
            "tables_count": tables_count,
            "total_chars": total_chars,
        }
        logger.info(f"[PDF-CONVERSION-VALIDATION] DOCX validated successfully: {stats}")
        return stats

    @classmethod
    def convert_pdf_file_to_docx(
        cls,
        pdf_path: str,
        output_docx_path: Optional[str] = None,
        preferred_provider: Optional[str] = None,
    ) -> Tuple[str, bytes, Dict[str, Any]]:
        """
        Converts a PDF file on disk to a high-fidelity DOCX file using the configured
        provider (e.g. CloudConvert) with automatic fallback to the internal converter (pdf2docx).
        """
        if not os.path.exists(pdf_path):
            raise ValueError("Source PDF file does not exist.")

        if not output_docx_path:
            base, _ = os.path.splitext(pdf_path)
            output_docx_path = f"{base}_working.docx"

        os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)

        conv_result, qgate_report = pdf_converter_registry.convert_with_quality_gate(
            pdf_path=pdf_path,
            output_docx_path=output_docx_path,
            preferred_provider=preferred_provider,
        )

        if not conv_result.success:
            logger.error(f"[PDF-TO-DOCX] All converter providers failed for '{pdf_path}': {conv_result.error}")
            raise ValueError(CONVERSION_ERROR_MESSAGE)

        # Validate the converted file
        stats = cls.validate_converted_docx(output_docx_path)
        stats["provider"] = conv_result.provider
        stats["fallback_used"] = conv_result.fallback_used
        stats["timing"] = conv_result.timing.model_dump()
        stats["quality_gate"] = qgate_report

        docx_bytes = conv_result.docx_bytes
        if not docx_bytes:
            with open(output_docx_path, "rb") as f:
                docx_bytes = f.read()

        return output_docx_path, docx_bytes, stats

    @classmethod
    def convert_pdf_bytes_to_docx(
        cls,
        pdf_bytes: bytes,
        output_docx_path: Optional[str] = None,
        original_filename: str = "template.pdf",
        preferred_provider: Optional[str] = None,
    ) -> Tuple[str, bytes, Dict[str, Any]]:
        """
        Converts PDF bytes into a validated DOCX file on disk and returns bytes.
        """
        if not pdf_bytes or len(pdf_bytes) == 0:
            raise ValueError("Uploaded PDF bytes are empty.")

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as temp_pdf:
            temp_pdf.write(pdf_bytes)
            temp_pdf_path = temp_pdf.name

        try:
            if not output_docx_path:
                stem = os.path.splitext(os.path.basename(original_filename))[0]
                with tempfile.NamedTemporaryFile(prefix=f"{stem}_", suffix=".docx", delete=False) as temp_docx:
                    output_docx_path = temp_docx.name

            return cls.convert_pdf_file_to_docx(
                temp_pdf_path,
                output_docx_path,
                preferred_provider=preferred_provider,
            )
        finally:
            if os.path.exists(temp_pdf_path):
                try:
                    os.remove(temp_pdf_path)
                except Exception:
                    pass


pdf_to_docx_service = PdfToDocxService()
