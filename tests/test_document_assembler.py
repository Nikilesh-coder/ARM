"""
Test Document Assembler Engine
"""

import os
import tempfile
import docx
from packages.template_intelligence.schema import TemplateSchema, TypographyRules, HeadingStyle
from packages.document_engine.engine import DocumentAssembler


def test_document_assembler_generates_docx():
    schema = TemplateSchema(
        template_id="tpl_gen_01",
        template_name="Generated Report Spec",
        typography=TypographyRules(
            default_font="Calibri",
            default_size_pt=11.5,
            headings={
                "heading_1": HeadingStyle(level=1, font_name="Calibri", size_pt=16.0, bold=True, all_caps=True),
                "heading_2": HeadingStyle(level=2, font_name="Calibri", size_pt=13.0, bold=True)
            }
        )
    )

    assembler = DocumentAssembler(template_schema=schema)
    assembler.add_heading("Chapter 1: System Overview", level=1)
    assembler.add_body_paragraph("ReportForge AI deterministic document assembly engine validates compliance.")
    assembler.add_heading("1.1 Architecture Components", level=2)
    assembler.add_table(
        headers=["Module", "Function", "Engine"],
        rows=[
            ["Template Intelligence", "OpenXML Extraction", "Deterministic Python AST"],
            ["AI Content", "Structured Section Synthesis", "Gemini Provider"],
            ["Document Assembler", "Style & Layout Synthesis", "python-docx"]
        ],
        caption="Table 1.1: System Architecture Modules"
    )

    temp_dir = tempfile.gettempdir()
    out_docx = os.path.join(temp_dir, "test_generated_report.docx")
    assembler.save(out_docx)

    assert os.path.exists(out_docx)
    # Read back with docx
    read_doc = docx.Document(out_docx)
    assert len(read_doc.paragraphs) > 0
    assert len(read_doc.tables) == 1
    assert len(read_doc.tables[0].rows) == 4  # 1 header + 3 data rows

    if os.path.exists(out_docx):
        os.remove(out_docx)
