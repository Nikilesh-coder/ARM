"""
ARM Template Preservation Rules Test Suite
Validates strict compliance with ARM Core Requirement:
- College Template is the Absolute Source of Truth
- Exact page dimensions & custom geometry preservation (no forced A4/Letter conversions)
- Strict distinction between FIXED and REPLACEABLE elements
- College logo and fixed shapes protection
- Section-by-section content limits and clamping
- In-place text and image replacement without document reconstruction
"""

import os
import io
import pytest
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

from apps.api.main import app
from packages.replacement_engine.base import (
    FIXED_ELEMENTS,
    REPLACEABLE_ELEMENTS,
    ReplacementState,
)
from packages.replacement_engine.docx_engine import DocxTemplateReplacementEngine

client = TestClient(app)


def test_fixed_vs_replaceable_elements_definition():
    """Verifies that FIXED and REPLACEABLE elements are strictly delineated."""
    assert "college_logo" in FIXED_ELEMENTS
    assert "college_name" in FIXED_ELEMENTS
    assert "university_logo" in FIXED_ELEMENTS
    assert "department_logo" in FIXED_ELEMENTS
    assert "page_border" in FIXED_ELEMENTS
    assert "header" in FIXED_ELEMENTS
    assert "footer" in FIXED_ELEMENTS
    assert "background" in FIXED_ELEMENTS
    assert "watermark" in FIXED_ELEMENTS
    assert "fixed_design_image" in FIXED_ELEMENTS
    assert "fixed_shape" in FIXED_ELEMENTS

    # Ensure no overlap
    overlap = FIXED_ELEMENTS.intersection(REPLACEABLE_ELEMENTS)
    assert len(overlap) == 0, f"Overlap between fixed and replaceable elements: {overlap}"

    # Verify canonical replaceable fields
    assert "project_title" in REPLACEABLE_ELEMENTS
    assert "student_name" in REPLACEABLE_ELEMENTS
    assert "introduction" in REPLACEABLE_ELEMENTS
    assert "methodology" in REPLACEABLE_ELEMENTS
    assert "results" in REPLACEABLE_ELEMENTS
    assert "conclusion" in REPLACEABLE_ELEMENTS
    assert "image_1" in REPLACEABLE_ELEMENTS


def test_exact_page_dimensions_preservation(tmp_path):
    """
    CRITICAL RULE: ARM must preserve the exact original page size and dimensions.
    Do NOT automatically change the document to A4, Letter, Legal or any default size.
    If original document uses custom dimensions, preserve those custom dimensions.
    """
    template_path = str(tmp_path / "custom_geometry_template.docx")
    output_path = str(tmp_path / "custom_geometry_output.docx")

    # 1. Create a template with non-standard CUSTOM dimensions and margins
    doc = docx.Document()
    sec = doc.sections[0]
    custom_width = Inches(7.25)   # Custom width (neither A4 nor Letter)
    custom_height = Inches(9.75)  # Custom height
    custom_left = Inches(1.35)
    custom_right = Inches(0.85)
    custom_top = Inches(1.15)
    custom_bottom = Inches(1.20)

    sec.page_width = custom_width
    sec.page_height = custom_height
    sec.left_margin = custom_left
    sec.right_margin = custom_right
    sec.top_margin = custom_top
    sec.bottom_margin = custom_bottom

    doc.add_paragraph("COLLEGE OF ENGINEERING (FIXED NAME)")
    doc.add_paragraph("Title: {{PROJECT_TITLE}}")
    doc.add_paragraph("Author: {{STUDENT_NAME}}")
    doc.add_paragraph("{{INTRODUCTION}}")
    doc.save(template_path)

    # 2. Execute in-place replacement
    engine = DocxTemplateReplacementEngine()
    result = engine.replace(
        template_path=template_path,
        field_values={
            "project_title": "AI Autonomous Robotics System",
            "student_name": "Test Student",
            "introduction": "This research analyzes edge AI control loops on constrained microcontrollers.",
        },
        output_path=output_path,
    )

    assert result.success is True

    # 3. Verify output document preserves EXACT custom geometry down to the inch
    out_doc = docx.Document(output_path)
    out_sec = out_doc.sections[0]

    assert abs(out_sec.page_width.inches - custom_width.inches) < 0.01, "Page width was modified!"
    assert abs(out_sec.page_height.inches - custom_height.inches) < 0.01, "Page height was modified!"
    assert abs(out_sec.left_margin.inches - custom_left.inches) < 0.01, "Left margin was modified!"
    assert abs(out_sec.right_margin.inches - custom_right.inches) < 0.01, "Right margin was modified!"
    assert abs(out_sec.top_margin.inches - custom_top.inches) < 0.01, "Top margin was modified!"
    assert abs(out_sec.bottom_margin.inches - custom_bottom.inches) < 0.01, "Bottom margin was modified!"


def test_fixed_elements_protection_and_college_logo(tmp_path):
    """
    Verifies that FIXED elements (college_logo, college_name, header, footer, page borders)
    are strictly protected from alteration, even if passed in input data.
    """
    template_path = str(tmp_path / "fixed_elements_template.docx")
    output_path = str(tmp_path / "fixed_elements_output.docx")

    doc = docx.Document()
    # Fixed College Title paragraph
    p_college = doc.add_paragraph("SRI VENKATESWARA COLLEGE OF ENGINEERING")
    p_college.runs[0].bold = True
    p_college.runs[0].font.size = Pt(16)
    p_college.runs[0].font.color.rgb = RGBColor(180, 0, 0)

    # Replaceable Project Title
    p_proj = doc.add_paragraph("PROJECT: {{PROJECT_TITLE}}")

    # Replaceable Image placeholder
    doc.add_paragraph("FIGURE 1: ARCHITECTURE")
    doc.add_paragraph("{{IMAGE_1}}")

    doc.save(template_path)

    # Create dummy PNG image asset
    from PIL import Image
    img_path = str(tmp_path / "real_chart.png")
    im = Image.new("RGB", (800, 500), color="#1e3a8a")
    im.save(img_path, dpi=(300, 300))

    # Attempt to pass a replacement for fixed elements (should be ignored)
    engine = DocxTemplateReplacementEngine()
    result = engine.replace(
        template_path=template_path,
        field_values={
            "college_name": "MALICIOUS OVERWRITE COLLEGE",
            "college_logo": "malicious_logo.png",
            "project_title": "SMART EMBEDDED ATTENDANCE SYSTEM",
        },
        image_assets={"image_1": img_path},
        output_path=output_path,
    )

    assert result.success is True
    out_doc = docx.Document(output_path)
    all_text = " ".join([p.text for p in out_doc.paragraphs])

    # Fixed college name must be intact and NOT overwritten
    assert "SRI VENKATESWARA COLLEGE OF ENGINEERING" in all_text
    assert "MALICIOUS OVERWRITE COLLEGE" not in all_text

    # Replaceable element was replaced
    assert "SMART EMBEDDED ATTENDANCE SYSTEM" in all_text
    assert "{{IMAGE_1}}" not in all_text


def test_content_length_limits_prevent_layout_destruction(tmp_path):
    """
    Verifies that AI content exceeding field limits is clamped / shortened
    to protect the template layout rather than expanding uncontrollably.
    """
    template_path = str(tmp_path / "limits_template.docx")
    output_path = str(tmp_path / "limits_output.docx")

    doc = docx.Document()
    doc.add_paragraph("CHAPTER 1: INTRODUCTION")
    doc.add_paragraph("{{INTRODUCTION}}")
    doc.save(template_path)

    # Generate an excessively long text (1,500 words) where max_words is 800
    excessive_text = " ".join(["algorithm" for _ in range(1500)])

    engine = DocxTemplateReplacementEngine()
    result = engine.replace(
        template_path=template_path,
        field_values={"introduction": excessive_text},
        output_path=output_path,
    )

    assert result.success is True
    out_doc = docx.Document(output_path)
    intro_p = [p for p in out_doc.paragraphs if "algorithm" in p.text][0]
    word_count = len(intro_p.text.split())

    # Introduction limit in CANONICAL_FIELDS is 800 words
    assert word_count <= 800, f"Expected <= 800 words, got {word_count}"


def test_reports_generate_endpoint_uses_replacement_engine():
    """
    Verifies that POST /api/v1/reports/generate uses the template-preserving
    replacement engine, returns valid download URLs, and does not build from scratch.
    """
    payload = {
        "title": "Autonomous Drone Navigation via Edge Compute",
        "project_type": "Capstone Project",
        "student_name": "Kasireddy Charishma",
        "roll_no": "25BFA02L13",
        "guide_name": "Dr. R. Sireesha Ph.D.",
        "department": "Department of Computer Science & Engineering",
        "institution": "Sri Venkateswara College of Engineering",
        "problem_statement": "Real-time trajectory planning with obstacle avoidance.",
    }

    res = client.post("/api/v1/reports/generate", json=payload)
    assert res.status_code == 200
    data = res.json()

    assert data["status"] == "completed"
    assert data["state"] == "COMPLETED"
    assert "report_id" in data
    assert "download_url" in data
    assert "pdf_download_url" in data
    assert data["fields_replaced"] >= 5
    assert len(data["preview_sections"]) > 0

    # Verify download works and returns genuine binary docx
    rep_id = data["report_id"]
    res_dl = client.get(f"/api/v1/reports/{rep_id}/download")
    assert res_dl.status_code == 200
    assert len(res_dl.content) > 10000
    assert res_dl.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
