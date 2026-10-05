"""
ARM Stage 2 - Comprehensive Automated Verification Suite
Validates:
1. College Template Upload (DOCX & PDF)
2. File Validation (Empty, Oversized, Unsupported, Corrupted)
3. Private Storage & Raw Byte Preservation (SHA-256)
4. Database Registration & Statuses ('uploaded', 'pending')
5. Safe Template Replacement (Preserving old original file)
6. Generation Job Creation (Analyze Template action queues job)
7. Multi-Tenant Authorization & IDOR Defense (User B cannot access User A's template)
"""

import io
import os
import uuid
import hashlib
import docx
import pypdf
import pytest
from fastapi.testclient import TestClient
from apps.api.main import app
from apps.api.core.auth import get_current_user_optional
from apps.api.core.database import db_manager
from apps.api.services.storage import get_storage_provider
from apps.api.core.config import settings

client = TestClient(app)

def create_sample_docx(text="ARM Academic Template Specification") -> io.BytesIO:
    doc = docx.Document()
    doc.add_heading("SRI VENKATESWARA COLLEGE OF ENGINEERING", level=0)
    doc.add_paragraph("Department of Computer Science & Engineering")
    doc.add_paragraph(text)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream

def create_sample_pdf(text="ARM Sample College PDF Template") -> io.BytesIO:
    from reportlab.pdfgen import canvas
    stream = io.BytesIO()
    c = canvas.Canvas(stream)
    c.drawString(100, 750, "COLLEGE OF ENGINEERING & TECHNOLOGY")
    c.drawString(100, 730, text)
    c.showPage()
    c.save()
    stream.seek(0)
    return stream


class TestStage2TemplateIngestion:
    user_a = {"id": "b64e4002-066c-4050-ba4e-bf0518369563", "email": "student@university.edu"}
    user_b = {"id": "1b2f3346-ca7f-4b38-916e-a3cb62a6e9cd", "email": "2024csm.r260@svce.edu.in"}
    project_a_id = None

    @pytest.fixture(autouse=True)
    def setup_project(self):
        # Create project for User A
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        res = client.post("/api/v1/projects", json={
            "title": "Autonomous Systems Major Project",
            "project_type": "capstone",
            "academic_year": "2026-2027",
            "guide_name": "Dr. R. Ramanujan"
        })
        assert res.status_code == 200
        self.project_a_id = res.json()["id"]
        yield
        # Clean up
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        client.delete(f"/api/v1/projects/{self.project_a_id}")
        app.dependency_overrides.pop(get_current_user_optional, None)

    def test_upload_valid_docx(self):
        """Verify successful upload of a valid DOCX college template."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        docx_stream = create_sample_docx("Valid College Formatting Specification")
        raw_bytes = docx_stream.getvalue()
        expected_hash = hashlib.sha256(raw_bytes).hexdigest()

        files = {"file": ("college_format.docx", docx_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 200
        data = res.json()

        assert data["name"] == "college_format.docx"
        assert data["file_type"] == "docx"
        assert data["status"] == "uploaded"
        assert data["analysis_status"] == "pending"
        assert data["project_id"] == self.project_a_id

        # Verify storage & byte preservation
        storage_path = data["original_file_path"]
        assert storage_path.startswith(f"users/{self.user_a['id']}/projects/{self.project_a_id}/templates/")
        assert storage_path.endswith("/original.docx")

        storage = get_storage_provider()
        stored_bytes = storage.download_file(settings.storage.bucket_templates, storage_path)
        assert stored_bytes == raw_bytes, "Original DOCX bytes were modified during ingestion!"
        assert hashlib.sha256(stored_bytes).hexdigest() == expected_hash

    def test_upload_valid_pdf(self):
        """Verify successful upload of a valid PDF college template."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        pdf_stream = create_sample_pdf("Valid PDF College Guidelines")
        raw_bytes = pdf_stream.getvalue()

        files = {"file": ("institutional_guideline.pdf", pdf_stream, "application/pdf")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 200
        data = res.json()

        assert data["name"] == "institutional_guideline.pdf"
        assert data["file_type"] == "pdf"
        assert data["status"] == "uploaded"
        assert data["analysis_status"] == "pending"
        assert data["original_file_path"].endswith("/original.pdf")

        # Verify raw bytes in private storage
        storage = get_storage_provider()
        stored_bytes = storage.download_file(settings.storage.bucket_templates, data["original_file_path"])
        assert stored_bytes == raw_bytes, "Original PDF bytes were modified during ingestion!"

    def test_file_validation_unsupported_extension(self):
        """Reject files with unsupported extensions (e.g. .txt, .jpg)."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        fake_stream = io.BytesIO(b"Plain text format specification")
        files = {"file": ("guidelines.txt", fake_stream, "text/plain")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 400
        assert "Unsupported file type" in res.json()["detail"]

    def test_file_validation_empty_file(self):
        """Reject 0-byte empty files."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        empty_stream = io.BytesIO(b"")
        files = {"file": ("empty_template.docx", empty_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 400
        assert "empty" in res.json()["detail"].lower()

    def test_file_validation_oversized_file(self):
        """Reject files exceeding 25MB threshold."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        # 26MB fake payload
        oversized_bytes = b"PK\x03\x04" + b"0" * (26 * 1024 * 1024)
        files = {"file": ("massive_template.docx", io.BytesIO(oversized_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 400
        assert "large" in res.json()["detail"].lower()

    def test_file_validation_corrupted_docx(self):
        """Reject corrupted DOCX files that lack Zip or OpenXML manifest."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        corrupted_bytes = b"This is not a zip archive"
        files = {"file": ("corrupt.docx", io.BytesIO(corrupted_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 400
        assert "could not read" in res.json()["detail"].lower()

    def test_file_validation_corrupted_pdf(self):
        """Reject corrupted PDF files that fail PDF parsing."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        corrupted_pdf_bytes = b"%PDF-1.4\nCorrupted content without xref or trailer"
        files = {"file": ("broken.pdf", io.BytesIO(corrupted_pdf_bytes), "application/pdf")}
        res = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res.status_code == 400
        assert "could not read" in res.json()["detail"].lower()

    def test_authorization_and_idor_protection(self):
        """Ensure User B cannot upload to, read, delete, or trigger analysis on User A's project template."""
        # 1. User A uploads template
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        docx_stream = create_sample_docx("User A Proprietary Template")
        files = {"file": ("user_a_template.docx", docx_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res_a = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files)
        assert res_a.status_code == 200
        template_id = res_a.json()["template_id"]

        # 2. Switch to User B
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_b

        # User B attempts upload into User A's project
        files_b = {"file": ("malicious.docx", create_sample_docx("B"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res_b_upload = client.post(f"/api/v1/projects/{self.project_a_id}/templates/upload", files=files_b)
        assert res_b_upload.status_code == 403, f"IDOR VULN: User B uploaded to User A's project! Status: {res_b_upload.status_code}"

        # User B attempts to read User A's project template
        res_b_get = client.get(f"/api/v1/projects/{self.project_a_id}/template")
        assert res_b_get.status_code == 403, f"IDOR VULN: User B read User A's template! Status: {res_b_get.status_code}"

        # User B attempts to trigger analysis on User A's template
        res_b_analyze = client.post(f"/api/v1/projects/{self.project_a_id}/templates/{template_id}/analyze")
        assert res_b_analyze.status_code == 403, f"IDOR VULN: User B triggered analysis on User A's template!"

        # User B attempts to delete User A's template
        res_b_del = client.delete(f"/api/v1/projects/{self.project_a_id}/templates/{template_id}")
        assert res_b_del.status_code == 403, f"IDOR VULN: User B deleted User A's template!"

    def test_safe_template_replacement(self):
        """Verify replacing a template preserves the previous original file and creates a new record."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a
        storage = get_storage_provider()

        # 1. Upload initial template
        stream_1 = create_sample_docx("First Template Edition")
        bytes_1 = stream_1.getvalue()
        res_1 = client.post(
            f"/api/v1/projects/{self.project_a_id}/templates/upload",
            files={"file": ("edition_1.docx", stream_1, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_1.status_code == 200
        path_1 = res_1.json()["original_file_path"]
        id_1 = res_1.json()["template_id"]

        # 2. Upload replacement template
        stream_2 = create_sample_docx("Second Template Edition (Updated Margins)")
        bytes_2 = stream_2.getvalue()
        res_2 = client.post(
            f"/api/v1/projects/{self.project_a_id}/templates/upload",
            files={"file": ("edition_2.docx", stream_2, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_2.status_code == 200
        path_2 = res_2.json()["original_file_path"]
        id_2 = res_2.json()["template_id"]

        assert id_1 != id_2, "Replacement should create a new template record ID!"
        assert path_1 != path_2, "Replacement should have unique storage paths!"

        # Verify BOTH original files exist in storage (Original 1 is NOT destroyed)
        stored_1 = storage.download_file(settings.storage.bucket_templates, path_1)
        stored_2 = storage.download_file(settings.storage.bucket_templates, path_2)
        assert stored_1 == bytes_1, "Previous template file was altered or destroyed!"
        assert stored_2 == bytes_2, "New template file was not properly saved!"

        # Verify GET active template returns the second template
        res_active = client.get(f"/api/v1/projects/{self.project_a_id}/template")
        assert res_active.status_code == 200
        assert res_active.json()["id"] == id_2

    def test_analyze_template_action_queues_generation_job(self):
        """Verify 'Analyze Template' action creates real generation_job with status 'queued'."""
        app.dependency_overrides[get_current_user_optional] = lambda: self.user_a

        # Upload template
        stream = create_sample_docx("Template for Job Queue Test")
        res_upload = client.post(
            f"/api/v1/projects/{self.project_a_id}/templates/upload",
            files={"file": ("report_format.docx", stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_upload.status_code == 200
        template_id = res_upload.json()["template_id"]

        # Trigger Analyze Template action
        res_analyze = client.post(f"/api/v1/projects/{self.project_a_id}/templates/{template_id}/analyze")
        assert res_analyze.status_code == 200
        data = res_analyze.json()

        assert "job_id" in data
        assert data["status"] == "queued"
        assert data["template_id"] == template_id
        assert data["project_id"] == self.project_a_id
        job_id = data["job_id"]

        # Verify generation_jobs record in database
        client_db = db_manager.client
        if client_db:
            job_res = client_db.table("generation_jobs").select("*").eq("id", job_id).execute()
            assert job_res.data and len(job_res.data) > 0
            job = job_res.data[0]
            assert job["job_type"] == "template_analysis"
            assert job["status"] == "queued"
            assert job["progress"] == 0
            assert job["current_step"] == "Analysis queued"
            assert job["project_id"] == self.project_a_id

            # Verify template status transitioned to analyzing
            t_res = client_db.table("templates").select("status, analysis_status").eq("id", template_id).execute()
            assert t_res.data[0]["status"] == "analyzing"
            assert t_res.data[0]["analysis_status"] == "pending"
