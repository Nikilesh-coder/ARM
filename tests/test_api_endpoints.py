"""
Test API Endpoints for Projects, Templates, and Reports
"""

import io
import docx
from fastapi.testclient import TestClient
from apps.api.main import app

client = TestClient(app)


def test_create_and_get_project():
    payload = {
        "title": "Autonomous Academic Documentation Pipeline",
        "project_type": "capstone",
        "academic_year": "2026-2027",
        "guide_name": "Dr. Sarah Jenkins",
        "abstract_summary": "Automated report synthesis respecting institutional guidelines.",
        "tech_stack": ["FastAPI", "Next.js", "Python-docx", "Supabase"]
    }
    res = client.post("/api/v1/projects", json=payload)
    assert res.status_code == 200
    data = res.json()
    project_id = data["id"]
    assert data["title"] == payload["title"]
    assert data["status"] == "draft"

    # Get project
    res_get = client.get(f"/api/v1/projects/{project_id}")
    assert res_get.status_code == 200
    assert res_get.json()["id"] == project_id


def test_template_upload_and_schema_endpoint():
    # Build sample docx in memory
    doc = docx.Document()
    doc.add_paragraph("CERTIFICATE OF AUTHENTICITY")
    doc.add_paragraph("CHAPTER 1 INTRODUCTION")
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    # Upload
    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    files = {"file": ("sample_college_template.docx", file_stream, docx_mime)}
    res = client.post("/api/v1/templates/upload", files=files)
    assert res.status_code == 200
    data = res.json()
    template_id = data["template_id"]
    assert data["status"] == "analyzed"

    # Fetch schema
    res_schema = client.get(f"/api/v1/templates/{template_id}/schema")
    assert res_schema.status_code == 200
    schema_data = res_schema.json()
    assert schema_data["template_id"] == template_id
    assert "geometry" in schema_data
    assert "typography" in schema_data


def test_create_report_plan():
    payload = {
        "project_id": "proj_123",
        "template_id": "tpl_123",
        "report_title": "Final Capstone Documentation",
        "report_type": "capstone"
    }
    res = client.post("/api/v1/reports/plan", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == payload["report_title"]
    assert len(data["sections"]) > 0
    assert data["sections"][0]["section_key"] == "title_page"
