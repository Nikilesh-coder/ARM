"""
Test Template Parser & Schema Extraction
"""

import os
import tempfile
import docx
from docx.shared import Inches
from packages.template_intelligence.parser import TemplateParser
from packages.template_intelligence.schema import TemplateSchema


def test_template_parser_extracts_schema():
    # 1. Create a dummy college template DOCX
    temp_dir = tempfile.gettempdir()
    sample_doc_path = os.path.join(temp_dir, "test_college_template.docx")

    doc = docx.Document()
    # Set margins
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)

    # Add sample academic sections
    doc.add_paragraph("A PROJECT REPORT ON ADVANCED AI SYSTEMS")
    doc.add_paragraph("CERTIFICATE OF APPROVAL")
    doc.add_paragraph("CANDIDATE'S DECLARATION")
    doc.add_paragraph("ACKNOWLEDGEMENTS")
    doc.add_paragraph("ABSTRACT")
    doc.add_paragraph("TABLE OF CONTENTS")
    doc.add_heading("CHAPTER 1 INTRODUCTION", level=1)
    doc.add_paragraph("This is the introduction text formatted properly.")
    doc.add_heading("1.1 Background", level=2)
    doc.add_paragraph("Background details.")
    doc.add_heading("REFERENCES", level=1)

    doc.save(sample_doc_path)

    # 2. Run Parser
    parser = TemplateParser(sample_doc_path)
    schema = parser.parse(template_id="tpl_test_01", template_name="Test College Template")

    # 3. Assertions
    assert isinstance(schema, TemplateSchema)
    assert schema.template_id == "tpl_test_01"
    assert schema.geometry.page_size == "A4" or schema.geometry.page_size == "Letter"
    assert schema.geometry.margins.left_mm > 30.0  # 1.25 in = 31.75 mm

    # Check detected sections
    section_types = [s.section_type for s in schema.sections]
    assert "certificate" in section_types
    assert "declaration" in section_types
    assert "acknowledgement" in section_types
    assert "abstract" in section_types
    assert "introduction" in section_types
    assert "references" in section_types

    # Clean up
    if os.path.exists(sample_doc_path):
        os.remove(sample_doc_path)
