"""
ARM Isolated Diagnostic Test: RAW TEMPLATE DOWNLOAD
Tests ONLY:
UPLOAD DOCX
→ STORE RAW FILE
→ DOWNLOAD SAME RAW FILE

Guarantees:
- Zero report generation
- Zero AI
- Zero replacement engine
- Zero PackageLevelDocxEngine
- Zero python-docx re-serialization
- Exact SHA-256 byte-for-byte equality between uploaded and downloaded files.
"""

import os
import io
import hashlib
import pytest
from fastapi.testclient import TestClient
from PIL import Image as PILImage
import docx
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from apps.api.main import app

client = TestClient(app)


@pytest.fixture
def sample_college_raw_docx(tmp_path):
    """Creates a sample college DOCX with borders, logo, and styles."""
    doc_path = str(tmp_path / "original_college_test_template.docx")
    doc = docx.Document()

    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.0)

    # Page borders
    sectPr = sec._sectPr
    bdr_xml = (
        f'<w:pgBorders {nsdecls("w")}>'
        f'<w:top w:val="double" w:sz="18" w:space="24" w:color="002060"/>'
        f'<w:left w:val="double" w:sz="18" w:space="24" w:color="002060"/>'
        f'<w:bottom w:val="double" w:sz="18" w:space="24" w:color="002060"/>'
        f'<w:right w:val="double" w:sz="18" w:space="24" w:color="002060"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(parse_xml(bdr_xml))

    # Logo
    logo = PILImage.new("RGB", (70, 70), color=(10, 50, 130))
    buf = io.BytesIO()
    logo.save(buf, format="PNG")
    buf.seek(0)
    p = doc.add_paragraph()
    p.add_run().add_picture(buf, width=Inches(1.0))

    # Text
    p2 = doc.add_paragraph("SRI VENKATESWARA COLLEGE OF ENGINEERING")
    p2.runs[0].bold = True
    doc.add_paragraph("PROJECT TITLE")
    doc.add_paragraph("CANDIDATE: K. CHARISHMA (25BFA02L13)")

    doc.save(doc_path)
    return doc_path


def test_upload_store_and_raw_download_identical_bytes(sample_college_raw_docx):
    """
    1. Read original college .docx bytes & compute SHA-256 and size.
    2. Upload file to /api/v1/templates/upload-custom.
    3. Call GET /api/v1/templates/{template_id}/raw-download.
    4. Assert downloaded bytes are 100% byte-for-byte identical (matching SHA-256 and size).
    """
    with open(sample_college_raw_docx, "rb") as f:
        original_bytes = f.read()

    original_size = len(original_bytes)
    original_sha256 = hashlib.sha256(original_bytes).hexdigest()

    # Upload
    upload_res = client.post(
        "/api/v1/templates/upload-custom",
        data={
            "name": "Original Raw College Template",
            "institution": "SVCE",
            "department": "EEE",
            "report_type": "Capstone",
        },
        files={
            "file": ("my_original_college.docx", original_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
    )

    assert upload_res.status_code == 200, f"Upload failed: {upload_res.text}"
    template_data = upload_res.json()["template"]
    template_id = template_data["id"]

    # Raw Download
    download_res = client.get(f"/api/v1/templates/{template_id}/raw-download")
    assert download_res.status_code == 200, f"Raw download failed: {download_res.text}"

    downloaded_bytes = download_res.content
    downloaded_size = len(downloaded_bytes)
    downloaded_sha256 = hashlib.sha256(downloaded_bytes).hexdigest()

    # PRINT EXACT DIAGNOSTIC LOG
    print("\n" + "=" * 60)
    print("RAW DOWNLOAD TEST:")
    print(f"Original SHA-256:   {original_sha256}")
    print(f"Downloaded SHA-256: {downloaded_sha256}")
    print(f"Same SHA-256:       {'YES' if original_sha256 == downloaded_sha256 else 'NO'}")
    print(f"Original size:     {original_size}")
    print(f"Downloaded size:   {downloaded_size}")
    print(f"Same size:         {'YES' if original_size == downloaded_size else 'NO'}")
    print("=" * 60 + "\n")

    # CRITICAL ASSERTIONS
    assert downloaded_size == original_size, f"Size altered! Orig={original_size}, Down={downloaded_size}"
    assert downloaded_sha256 == original_sha256, f"Hash altered! Orig={original_sha256}, Down={downloaded_sha256}"
    assert download_res.headers.get("X-Stored-SHA256") == original_sha256


def test_master_template_raw_download():
    """Validates raw download of the default Master College Template."""
    master_source = os.path.abspath(".storage/templates/master_college_template.docx")
    assert os.path.exists(master_source)

    with open(master_source, "rb") as f:
        master_bytes = f.read()

    master_sha256 = hashlib.sha256(master_bytes).hexdigest()
    master_size = len(master_bytes)

    res = client.get("/api/v1/templates/00000000-0000-0000-0000-000000000001/raw-download")
    assert res.status_code == 200
    down_bytes = res.content
    assert len(down_bytes) == master_size
    assert hashlib.sha256(down_bytes).hexdigest() == master_sha256


def test_download_specific_template_id_140ada3976a8():
    """
    Directly tests downloading the exact template ID 'tpl_cust_140ada3976a8' reported by user.
    Verifies disk lookup successfully finds the stored JAGADEESH PPT.docx file.
    """
    target_id = "tpl_cust_140ada3976a8"
    res = client.get(f"/api/v1/templates/{target_id}/raw-download")
    assert res.status_code == 200, f"Failed to download {target_id}: {res.text}"

    down_bytes = res.content
    assert len(down_bytes) == 558982
    expected_sha256 = "a0847c7ca9fe58783703b0200f8678cb2b4158d73da4a8165703ce21e42569c4"
    assert hashlib.sha256(down_bytes).hexdigest() == expected_sha256
    print(f"\n[DOWNLOAD VERIFIED: {target_id}] Size: {len(down_bytes)}, SHA256: {expected_sha256}")

