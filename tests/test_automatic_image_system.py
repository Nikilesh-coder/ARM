"""
Test Suite for ARM Automatic Image System
Tests:
- Requirement determination from project topic, section, and template requirements
- Execution of Automatic Image Pipeline (generation, validation, storage, ReportAsset creation)
- Verification of ReportAsset fields: asset_type, source, description, file_path, template_field, project_id, attribution/license
- Strict error reporting when unavailable/unconfigured external API is requested
- Template image mapping and DocumentAssembler image replacement
"""

import os
import io
import docx
from fastapi.testclient import TestClient
from apps.api.main import app
from packages.template_intelligence.schema import TemplateSchema, DocumentGeometry, PageMargins, TypographyRules, HeaderFooterRule
from packages.document_engine.engine import DocumentAssembler

client = TestClient(app)


def test_determine_image_requirements_endpoint():
    payload = {
        "project_title": "Autonomous Underwater Inspection Vehicle",
        "project_description": "An autonomous subsea vehicle for inspecting offshore pipeline integrity using acoustic sonar and computer vision.",
        "template_id": "00000000-0000-0000-0000-000000000001"
    }

    res = client.post("/api/v1/projects/proj_test_img_1/assets/determine-requirements", json=payload)
    assert res.status_code == 200
    data = res.json()

    assert data["total_image_fields"] >= 3
    req_fields = {r["template_field"]: r for r in data["requirements"]}

    # Verify image_1 is identified as architecture diagram in Chapter 2
    assert "image_1" in req_fields
    assert req_fields["image_1"]["asset_type"] == "architecture_diagram"
    assert "Methodology" in req_fields["image_1"]["section_key"]
    assert req_fields["image_1"]["target_dimensions"]["dpi"] == 300
    assert "Academic Open" in req_fields["image_1"]["license_type"]

    # Verify image_2 is identified as implementation setup in Chapter 3
    assert "image_2" in req_fields
    assert req_fields["image_2"]["asset_type"] == "implementation_setup"
    assert "Implementation" in req_fields["image_2"]["section_key"]

    # Verify image_3 is identified as empirical chart in Chapter 4
    assert "image_3" in req_fields
    assert req_fields["image_3"]["asset_type"] == "empirical_chart"
    assert "Results" in req_fields["image_3"]["section_key"]


def test_automatic_image_pipeline_execution_and_report_assets():
    project_id = "proj_auto_img_test_100"
    payload = {
        "project_id": project_id,
        "project_title": "Autonomous Solar Microgrid Controller",
        "project_description": "Edge computing embedded system with maximum power point tracking and load prediction algorithms.",
        "template_id": "00000000-0000-0000-0000-000000000001"
    }

    res = client.post(f"/api/v1/projects/{project_id}/automatic-images", json=payload)
    assert res.status_code == 200
    data = res.json()

    assert data["status"] == "completed"
    assert data["successful_images_count"] >= 3
    assert len(data["assets"]) >= 3

    # Check field mapping contains all 3 images
    mapping = data["field_mapping"]
    assert "image_1" in mapping
    assert "image_2" in mapping
    assert "image_3" in mapping

    # Check ReportAsset records
    assets_by_field = {a["template_field"]: a for a in data["assets"]}
    for f in ["image_1", "image_2", "image_3"]:
        asset = assets_by_field[f]
        assert asset["project_id"] == project_id
        assert asset["template_field"] == f
        assert asset["source"] in ("diagram_engine", "chart_engine", "generated")
        assert len(asset["description"]) > 0
        assert os.path.exists(asset["storage_path"])
        assert asset["file_size_bytes"] > 1000  # Valid non-empty binary
        assert "attribution" in asset and len(asset["attribution"]) > 0
        assert "license_info" in asset and len(asset["license_info"]) > 0

    # Verify GET /projects/{id}/assets lists these records
    res_list = client.get(f"/api/v1/projects/{project_id}/assets")
    assert res_list.status_code == 200
    list_data = res_list.json()
    assert len(list_data) >= 3


def test_unconfigured_external_api_reports_error_not_faked():
    project_id = "proj_api_err_test"
    payload = {
        "project_id": project_id,
        "project_title": "AI Robotic Swarm",
        "project_description": "Robotic swarm exploration.",
        "preferred_provider": "unsplash",  # No UNSPLASH_ACCESS_KEY configured in test env
        "target_fields": ["image_1"]
    }

    # If UNSPLASH_ACCESS_KEY is not set in env, it should clearly report error
    if not os.getenv("UNSPLASH_ACCESS_KEY"):
        res = client.post(f"/api/v1/projects/{project_id}/automatic-images", json=payload)
        assert res.status_code == 200
        data = res.json()
        assert data["failed_images_count"] >= 1
        step = data["step_results"][0]
        assert step["status"] == "failed"
        assert "UNSPLASH_ACCESS_KEY" in step["error"] or "unavailable" in step["error"]


def test_document_assembler_template_image_replacement(tmp_path):
    # 1. Create a sample docx containing {{IMAGE_1}} and {{IMAGE_3}}
    doc = docx.Document()
    doc.add_paragraph("CHAPTER 2: SYSTEM METHODOLOGY")
    doc.add_paragraph("Below is the architecture diagram for the proposed design:")
    doc.add_paragraph("{{IMAGE_1}}")
    doc.add_paragraph("CHAPTER 4: RESULTS")
    doc.add_paragraph("{{IMAGE_3}}")

    tpl_path = str(tmp_path / "template_test.docx")
    doc.save(tpl_path)

    # 2. Generate sample images via the pipeline
    project_id = "proj_doc_replace_test"
    payload = {
        "project_id": project_id,
        "project_title": "Realtime Pipeline Test",
        "project_description": "Testing automated document replacement with real generated images.",
    }
    res = client.post(f"/api/v1/projects/{project_id}/automatic-images", json=payload)
    assert res.status_code == 200
    pipeline_data = res.json()
    field_mapping = pipeline_data["field_mapping"]

    # 3. Assemble document with replacement
    schema = TemplateSchema(
        template_id="tpl_test_replace",
        template_name="Test Template",
        geometry=DocumentGeometry(
            width_mm=210.0, height_mm=297.0,
            margins=PageMargins(top_mm=25.4, bottom_mm=25.4, left_mm=25.4, right_mm=25.4)
        ),
        typography=TypographyRules(default_font="Times New Roman", default_size_pt=12.0),
        header_footer=HeaderFooterRule()
    )
    assembler = DocumentAssembler(template_schema=schema, master_template_path=tpl_path)

    # Replace template images
    assembler.replace_template_images(field_mapping)

    output_docx = str(tmp_path / "output_replaced.docx")
    assembler.save(output_docx)

    # 4. Verify output document contains the embedded images (drawings / inline pictures)
    reloaded_doc = docx.Document(output_docx)
    all_text = " ".join([p.text for p in reloaded_doc.paragraphs])
    assert "{{IMAGE_1}}" not in all_text, "Placeholder {{IMAGE_1}} was not replaced!"
    assert "{{IMAGE_3}}" not in all_text, "Placeholder {{IMAGE_3}} was not replaced!"

    # Check for embedded drawings in document body xml
    xml_str = reloaded_doc._element.xml
    assert "w:drawing" in xml_str or "graphicData" in xml_str, "Document does not contain embedded image drawings!"
