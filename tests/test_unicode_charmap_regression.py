"""
ARM Unicode Charmap Encoding Regression Test
============================================
Verifies that template loading, slide analysis, image prompt logging, and
package replacement succeed without 'charmap' / cp1252 UnicodeEncodeError
when encountering Unicode ligatures like \ufb01 ('fi'), \ufb02 ('fl'),
en-dash, em-dash, and smart quotes.

Requirements verified:
- Template loading and parsing succeeds with \ufb01 present.
- Logging and printing of slide matter containing \ufb01 does not crash on Windows.
- OpenXML document retains valid UTF-8 structures.
- Unicode characters are preserved safely (not destroyed or corrupted).
"""

import os
import sys
import tempfile
import pytest
import docx

# Ensure root workspace is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from apps.api.core.logging import setup_logging
from apps.api.services.gemini_visual_pipeline_service import GeminiVisualPipelineService
from packages.replacement_engine.package_engine import PackageLevelDocxEngine


def test_unicode_ligature_in_slide_pipeline():
    """
    Tests that slide visual generation and logging handles \ufb01 safely without charmap crash.
    """
    setup_logging()

    # Replicate exact text from the template's Slide 34 that triggered the failure
    problematic_matter = "Make the area safe \ufb01rst. Handle electrical \ufb01res correctly. Avoid short circuits."
    slide_heading = "EMERGENCY RESPONSE & SAFETY \ufb01RST"

    req = GeminiVisualPipelineService.compile_visual_requirement(
        project_title="AI Based Irrigation System",
        slide_heading=slide_heading,
        slide_matter=problematic_matter,
        slide_index=34,
    )

    assert "\ufb01" in req.slide_matter
    assert "\ufb01" in req.slide_heading

    # Test generation and logging
    out_dir = os.path.abspath(os.path.join("generated_images", "regression_test"))
    os.makedirs(out_dir, exist_ok=True)
    out_img = os.path.join(out_dir, "slide_34_unicode_test.png")

    # Should execute without raising UnicodeEncodeError: 'charmap' codec can't encode character '\ufb01'
    result = GeminiVisualPipelineService.generate_and_replace_slide_image(
        project_title="AI Based Irrigation System",
        slide_heading=slide_heading,
        slide_matter=problematic_matter,
        slide_index=34,
        out_file_path=out_img,
        report_id="unicode_regression_test",
    )

    assert result is not None
    print("\n[REGRESSION TEST PASS] Slide visual requirement successfully processed with \ufb01 ligature.")


def test_template_with_unicode_ligatures_preservation():
    """
    Creates a temporary DOCX with \ufb01 and other special Unicode characters,
    verifying that PackageLevelDocxEngine processes and preserves it cleanly.
    """
    temp_dir = tempfile.mkdtemp()
    test_docx_path = os.path.join(temp_dir, "template_with_ligatures.docx")
    out_docx_path = os.path.join(temp_dir, "output_with_ligatures.docx")

    # Create docx with problematic Unicode characters
    doc = docx.Document()
    doc.add_heading("PROJECT TITLE: {{PROJECT_TITLE}}", level=0)
    doc.add_paragraph("Emergency Protocol: Make the area safe \ufb01rst (\ufb01 = Latin Small Ligature FI).")
    doc.add_paragraph("Secondary Procedure: Observe the \ufb02ow (\ufb02 = Latin Small Ligature FL).")
    doc.add_paragraph("Punctuation Test: En-dash 2020\u20132026, Em-dash \u2014, and \u201csmart quotes\u201d.")
    doc.save(test_docx_path)

    # Perform package replacement
    rep_result = PackageLevelDocxEngine.replace(
        template_path=test_docx_path,
        field_values={"project_title": "AI Based Irrigation System"},
        output_path=out_docx_path,
        job_id="regression_job_unicode",
    )

    assert rep_result.success is True
    assert os.path.exists(out_docx_path)

    # Read back and verify Unicode characters are preserved exactly
    doc_out = docx.Document(out_docx_path)
    full_text = "\n".join(p.text for p in doc_out.paragraphs)

    assert "AI Based Irrigation System" in full_text
    assert "\ufb01" in full_text, "Ligature \ufb01 must be preserved in output document"
    assert "\ufb02" in full_text, "Ligature \ufb02 must be preserved in output document"
    assert "\u2014" in full_text, "Em-dash must be preserved in output document"
    assert "\u201c" in full_text, "Smart quotes must be preserved in output document"

    print("\n[REGRESSION TEST PASS] DOCX template with Unicode ligatures parsed and preserved with 100% integrity.")


if __name__ == "__main__":
    test_unicode_ligature_in_slide_pipeline()
    test_template_with_unicode_ligatures_preservation()
    print("\nALL UNICODE CHARMAP REGRESSION TESTS PASSED SUCCESSFULLY!")
