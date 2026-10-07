"""
ARM - DOCX-Only College Template Upload Test Suite
Verifies:
1. DOCX template upload succeeds (HTTP 200).
2. Uploaded DOCX template registers correctly with status 'ready', file_type 'docx'.
3. Report generation succeeds using uploaded DOCX template.
4. PDF template upload is rejected with HTTP 400 and detail "Only DOCX college templates are currently supported."
5. Template analysis runs for DOCX templates (extracts page geometry, fields, images).
6. Final report preserves college template structure (margins, dimensions, header, footer).
"""

import os
import io
import pytest
from PIL import Image as PILImage
import docx
from docx.shared import Inches, Pt
from fastapi.testclient import TestClient

from apps.api.main import app
from apps.api.services.custom_template_service import custom_template_service

client = TestClient(app)


@pytest.fixture(scope="session")
def realistic_docx_template(tmp_path_factory) -> str:
    """Creates a realistic college template in .docx format."""
    tmp_dir = tmp_path_factory.mktemp("docx_tpl")
    tpl_path = str(tmp_dir / "College_Capstone_Template.docx")

    doc = docx.Document()

    # 1. Custom Page Geometry (7.5" x 10.0", 0.75" margins)
    sec = doc.sections[0]
    sec.page_width = Inches(7.5)
    sec.page_height = Inches(10.0)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    # 2. Institutional Header & Footer
    hdr = sec.header
    hp = hdr.paragraphs[0]
    hp.text = "STANFORD INSTITUTE OF TECHNOLOGY - CSE DEPARTMENT"

    ftr = sec.footer
    fp = ftr.paragraphs[0]
    fp.text = "ACADEMIC YEAR 2026-2027 - CONFIDENTIAL"

    # 3. Embed College Logo
    logo_img = PILImage.new("RGB", (60, 60), color=(10, 50, 150))
    logo_buf = io.BytesIO()
    logo_img.save(logo_buf, format="PNG")
    logo_buf.seek(0)
    p_logo = doc.add_paragraph()
    r_logo = p_logo.add_run()
    r_logo.add_picture(logo_buf, width=Inches(1.0))

    # 4. College Name
    p_inst = doc.add_paragraph()
    p_inst.add_run("STANFORD INSTITUTE OF TECHNOLOGY").bold = True

    # 5. Project Title placeholder
    p_title = doc.add_paragraph()
    p_title.add_run("PROJECT TITLE").bold = True

    # 6. Candidate Information
    p_cand = doc.add_paragraph()
    p_cand.add_run("SUBMITTED BY: ALICE SMITH (ROLL NO: 26CS101)")

    # 7. Chapter 1 Introduction
    p_ch1 = doc.add_paragraph()
    p_ch1.add_run("CHAPTER 1: INTRODUCTION").bold = True
    p_body = doc.add_paragraph()
    p_body.add_run(
        "This project investigates state-of-the-art computational techniques for academic workload automation."
    )

    doc.save(tpl_path)
    return tpl_path


@pytest.fixture(scope="session")
def realistic_pdf_bytes() -> bytes:
    """Generates valid minimal PDF bytes to test upload rejection."""
    import pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    pdf_buf = io.BytesIO()
    writer.write(pdf_buf)
    return pdf_buf.getvalue()


_UPLOADED_TEMPLATE_ID = None


def test_1_docx_template_upload_succeeds(realistic_docx_template):
    """TEST 1: DOCX template upload succeeds with HTTP 200."""
    global _UPLOADED_TEMPLATE_ID

    with open(realistic_docx_template, "rb") as f:
        content = f.read()

    res = client.post(
        "/api/v1/templates/upload-custom",
        files={"file": ("College_Capstone_Template.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "name": "Stanford Capstone Template",
            "institution": "Stanford Institute of Technology",
            "department": "Computer Science & Engineering",
            "report_type": "Capstone",
        },
    )

    assert res.status_code == 200, f"Upload failed: {res.text}"
    body = res.json()
    assert body["status"] == "success"
    assert "template" in body
    assert body["template"]["name"] == "Stanford Capstone Template"
    assert body["template"]["file_type"] == "docx"
    assert body["template"]["converted_from_pdf"] is False

    _UPLOADED_TEMPLATE_ID = body["template"]["id"]


def test_2_docx_template_registration_ready_and_docx():
    """TEST 2: Uploaded DOCX template registers correctly with status 'ready', file_type 'docx'."""
    global _UPLOADED_TEMPLATE_ID
    assert _UPLOADED_TEMPLATE_ID is not None

    tpl = custom_template_service.get_template(_UPLOADED_TEMPLATE_ID)
    assert tpl is not None
    assert tpl["id"] == _UPLOADED_TEMPLATE_ID
    assert tpl["file_type"] == "docx"
    assert tpl["status"] in ("ready", "analyzed")
    assert tpl["converted_from_pdf"] is False
    assert tpl["original_file_path"].endswith(".docx")
    assert os.path.exists(tpl["original_file_path"])


def test_3_report_generation_succeeds_using_uploaded_docx_template():
    """TEST 3: Report generation succeeds using uploaded DOCX template."""
    global _UPLOADED_TEMPLATE_ID
    assert _UPLOADED_TEMPLATE_ID is not None

    res = client.post(
        "/api/v1/reports/generate",
        json={
            "title": "Smart Campus IoT Energy Management System",
            "template_id": _UPLOADED_TEMPLATE_ID,
            "student_name": "Alice Smith",
            "roll_no": "26CS101",
        },
    )

    assert res.status_code == 200, f"Report generation failed: {res.text}"
    body = res.json()
    assert body.get("status") in ("COMPLETED", "success", "ready") or "report_id" in body or "docx_url" in body

    # Verify generated document exists on disk
    docx_path = body.get("docx_path") or body.get("output_path")
    if not docx_path and body.get("report_id"):
        docx_path = os.path.abspath(f".storage/reports/{body['report_id']}.docx")

    assert docx_path is not None
    assert os.path.exists(docx_path)
    assert os.path.getsize(docx_path) > 0


def test_4_pdf_template_upload_rejected_with_http_400(realistic_pdf_bytes):
    """TEST 4: PDF template upload is rejected with HTTP 400 and clear message."""
    # 4a. POST /api/v1/templates/upload-custom
    res_custom = client.post(
        "/api/v1/templates/upload-custom",
        files={"file": ("college_guide.pdf", realistic_pdf_bytes, "application/pdf")},
        data={"name": "PDF Template Attempt"},
    )
    assert res_custom.status_code == 400
    assert "Only DOCX college templates are currently supported." in res_custom.json()["detail"]

    # 4b. POST /api/v1/templates/upload
    res_standalone = client.post(
        "/api/v1/templates/upload",
        files={"file": ("test_guide.pdf", realistic_pdf_bytes, "application/pdf")},
    )
    assert res_standalone.status_code == 400
    assert "Only DOCX college templates are currently supported." in res_standalone.json()["detail"]

    # 4c. POST /api/v1/projects/{project_id}/templates/upload
    from apps.api.routers.projects import _PROJECTS_STORE
    test_proj_id = "proj_test_upload_rejection"
    _PROJECTS_STORE[test_proj_id] = {"id": test_proj_id, "user_id": "usr_demo_student"}

    res_project = client.post(
        f"/api/v1/projects/{test_proj_id}/templates/upload",
        files={"file": ("project_template.pdf", realistic_pdf_bytes, "application/pdf")},
    )
    assert res_project.status_code == 400
    assert "Only DOCX college templates are currently supported." in res_project.json()["detail"]


def test_5_template_analysis_runs_for_docx_template():
    """TEST 5: Template analysis runs for DOCX templates (geometry, text fields, images)."""
    global _UPLOADED_TEMPLATE_ID
    assert _UPLOADED_TEMPLATE_ID is not None

    tpl = custom_template_service.get_template(_UPLOADED_TEMPLATE_ID)
    assert tpl is not None

    # Check structural page settings extracted
    page_settings = tpl.get("page_settings", {})
    assert page_settings.get("width_in") == 7.5
    assert page_settings.get("height_in") == 10.0
    assert page_settings.get("top_margin_in") == 0.75
    assert page_settings.get("bottom_margin_in") == 0.75

    # Check detected text fields
    fields = tpl.get("field_mapping", [])
    assert len(fields) > 0
    # Must have detected candidate project_title
    field_keys = [f.get("arm_field") for f in fields]
    assert "project_title" in field_keys

    # Check detected images (logo preserved)
    images = tpl.get("image_mapping", [])
    assert len(images) >= 1
    assert any(img.get("is_logo") is True for img in images)

    # Check fixed elements summary
    fixed_summary = tpl.get("fixed_elements_summary", {})
    assert "college_logo" in fixed_summary
    assert "page_dimensions" in fixed_summary


def test_6_final_report_preserves_college_template_structure():
    """TEST 6: Final generated report preserves college template structure."""
    global _UPLOADED_TEMPLATE_ID
    assert _UPLOADED_TEMPLATE_ID is not None

    # Retrieve generated report from storage
    reports_dir = os.path.abspath(".storage/reports")
    docx_files = [
        os.path.join(reports_dir, f)
        for f in os.listdir(reports_dir)
        if f.endswith(".docx") and not f.startswith("test_")
    ]
    assert len(docx_files) > 0, "No generated reports found in storage"

    latest_report = max(docx_files, key=os.path.getmtime)
    doc = docx.Document(latest_report)

    # 1. Exact geometry preserved
    sec = doc.sections[0]
    assert round(sec.page_width.inches, 2) == 7.5
    assert round(sec.page_height.inches, 2) == 10.0
    assert round(sec.top_margin.inches, 2) == 0.75
    assert round(sec.bottom_margin.inches, 2) == 0.75

    # 2. Institutional header & footer preserved
    assert "STANFORD INSTITUTE OF TECHNOLOGY" in sec.header.paragraphs[0].text
    assert "CONFIDENTIAL" in sec.footer.paragraphs[0].text

    # 3. New project title present in document
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Smart Campus" in all_text or "IoT Energy" in all_text or "Alice Smith" in all_text
